from pathlib import Path
from datetime import datetime

import pandas as pd
import matplotlib.pyplot as plt

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# 1. PATHS AND INPUT FILES
# --------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
PROCESSED = ROOT / "data" / "processed"
RAW = ROOT / "data" / "raw"

REPORTS.mkdir(parents=True, exist_ok=True)

summary_path = REPORTS / "quality_check_summary.csv"
anomaly_path = PROCESSED / "transactions_anomaly_flagged.csv"
raw_path = RAW / "transactions_messy.csv"

for path in [summary_path, anomaly_path, raw_path]:
    if not path.exists():
        raise FileNotFoundError(
            f"Required file not found: {path}"
        )

start_time = datetime.now()

# Read existing outputs from previous phases
summary = pd.read_csv(summary_path)
df = pd.read_csv(raw_path)
anomalies = pd.read_csv(anomaly_path)

print("Input files loaded successfully.")


# --------------------------------------------------
# 2. VALIDATE REQUIRED COLUMNS
# --------------------------------------------------

required_columns = {
    "transaction_id",
    "customer_id",
    "transaction_date",
    "amount",
    "email",
}

missing_columns = required_columns - set(df.columns)

if missing_columns:
    raise ValueError(
        f"Missing required columns in raw dataset: {missing_columns}"
    )

summary_required = {
    "check_name",
    "dimension",
    "severity",
    "total_rows",
    "failed_rows",
    "pass_rows",
}

missing_summary_columns = summary_required - set(summary.columns)

if missing_summary_columns:
    raise ValueError(
        "Missing columns in quality_check_summary.csv: "
        f"{missing_summary_columns}"
    )

if "anomaly_flag" not in anomalies.columns:
    raise ValueError(
        "The Phase 5 anomaly file must contain 'anomaly_flag'. "
        f"Available columns: {list(anomalies.columns)}"
    )


# --------------------------------------------------
# 3. PREPARE DATA
# --------------------------------------------------

df["transaction_date"] = pd.to_datetime(
    df["transaction_date"], errors="coerce"
)

amount = pd.to_numeric(df["amount"], errors="coerce")
email = df["email"].fillna("").astype(str).str.strip()

# Missing values in important fields
missing_columns_to_check = [
    "transaction_id",
    "customer_id",
    "transaction_date",
    "amount",
    "email",
]

missing_mask = df[missing_columns_to_check].isna().any(axis=1)

missing_rows = df.loc[missing_mask].copy()

# Duplicate transaction IDs
duplicate_mask = (
    df["transaction_id"].notna()
    & df["transaction_id"].duplicated(keep=False)
)

duplicate_rows = df.loc[duplicate_mask].copy()

# Invalid amounts and email formats
invalid_mask = (
    amount.isna()
    | (amount <= 0)
    | (
        email.ne("")
        & ~email.str.match(
            r"^[^@\s]+@[^@\s]+\.[^@\s]+$"
        )
    )
)

# Validate currency only if the column exists
if "currency" in df.columns:
    valid_currencies = ["INR", "USD", "EUR", "GBP"]

    invalid_currency = (
        df["currency"].notna()
        & ~df["currency"].isin(valid_currencies)
    )

    invalid_mask |= invalid_currency

invalid_rows = df.loc[invalid_mask].copy()


# --------------------------------------------------
# 4. PREPARE ANOMALY HIGHLIGHTS
# --------------------------------------------------

def is_flagged(value):
    """Safely interpret common boolean and numeric flag values."""
    if pd.isna(value):
        return False

    if isinstance(value, bool):
        return value

    return str(value).strip().lower() in {
        "1", "1.0", "true", "yes"
    }


flagged = anomalies.loc[
    anomalies["anomaly_flag"].apply(is_flagged)
].copy()

# Identify detector flags that can explain why a row was flagged
excluded_flags = {
    "anomaly_flag",
    "ground_truth_flag",
    "ground_truth_issue_flag",
    "ground_truth_anomaly_flag",
}

reason_columns = [
    column
    for column in flagged.columns
    if column.endswith("_flag")
    and column not in excluded_flags
]


def get_reasons(row):
    reasons = [
        column.replace("_flag", "").replace("_", " ").title()
        for column in reason_columns
        if is_flagged(row[column])
    ]

    if reasons:
        return ", ".join(reasons)

    return "Combined anomaly detector"


flagged["flag_reason"] = flagged.apply(
    get_reasons, axis=1
)

# Put transactions with more detector flags first
flagged["_reason_count"] = flagged[reason_columns].apply(
    lambda row: sum(is_flagged(value) for value in row),
    axis=1
) if reason_columns else 0

flagged = flagged.sort_values(
    "_reason_count", ascending=False
)

anomaly_columns = [
    column
    for column in [
        "transaction_id",
        "customer_id",
        "transaction_date",
        "amount",
        "currency",
        "anomaly_flag",
        "flag_reason",
    ]
    if column in flagged.columns
]

top_anomalies = flagged[anomaly_columns].head(10).copy()


# --------------------------------------------------
# 5. CALCULATE QUALITY SUMMARY
# --------------------------------------------------

summary["failure_percent"] = (
    summary["failed_rows"]
    / summary["total_rows"].replace(0, pd.NA)
    * 100
).fillna(0).round(2)

summary["pass_rate_percent"] = (
    summary["pass_rows"]
    / summary["total_rows"].replace(0, pd.NA)
    * 100
).fillna(0).round(2)

# Overall score:
# Prefer a row-level ground-truth issue flag if available.
# Otherwise report a score based on the selected checks below.
ground_truth_column = next(
    (
        column
        for column in [
            "ground_truth_issue_flag",
            "ground_truth_anomaly_flag",
            "ground_truth_flag",
        ]
        if column in anomalies.columns
    ),
    None,
)

if ground_truth_column is not None:
    issue_flags = anomalies[ground_truth_column].apply(is_flagged)

    total_rows = len(anomalies)
    clean_rows = total_rows - int(issue_flags.sum())

    score_basis = (
        f"Rows without {ground_truth_column}"
    )

else:
    # Use the union of the selected checks.
    # This is NOT a score based on all Phase 4 checks.
    issue_mask = (
        missing_mask
        | duplicate_mask
        | invalid_mask
    )

    total_rows = len(df)
    clean_rows = total_rows - int(issue_mask.sum())

    score_basis = (
        "Selected missing-field, duplicate-ID and invalid-value checks only"
    )

dq_score = (
    clean_rows / total_rows * 100
    if total_rows > 0
    else 0
)

# Group Phase 4 results by dimension
dimension_summary = (
    summary.groupby("dimension", dropna=False)["failed_rows"]
    .sum()
    .sort_values(ascending=False)
)

# Rank severity correctly: High first, then Medium, then Low
severity_order = {
    "High": 0,
    "Medium": 1,
    "Low": 2,
}

summary["_severity_rank"] = (
    summary["severity"]
    .map(severity_order)
    .fillna(3)
)

top_issues = (
    summary.sort_values(
        ["_severity_rank", "failed_rows"],
        ascending=[True, False],
    )
    .head(10)
    .copy()
)


# --------------------------------------------------
# 6. CREATE DIMENSION CHART
# --------------------------------------------------

chart_path = REPORTS / "dq_dimension_chart.png"

plt.figure(figsize=(9, 5))

dimension_summary.plot(kind="bar")

plt.title("Data Quality Failures by Dimension")
plt.xlabel("Data Quality Dimension")
plt.ylabel("Failed Check Rows")
plt.xticks(rotation=35, ha="right")
plt.tight_layout()

plt.savefig(chart_path, dpi=160)
plt.close()

print("Dimension chart created.")


# --------------------------------------------------
# 7. EXPORT EXCEL WORKBOOK
# --------------------------------------------------

excel_path = REPORTS / "data_quality_report.xlsx"

with pd.ExcelWriter(
    excel_path, engine="openpyxl"
) as writer:

    summary.drop(
        columns=["_severity_rank"], errors="ignore"
    ).to_excel(
        writer, sheet_name="Summary", index=False
    )

    missing_rows.to_excel(
        writer, sheet_name="Missing", index=False
    )

    duplicate_rows.to_excel(
        writer, sheet_name="Duplicates", index=False
    )

    invalid_rows.to_excel(
        writer, sheet_name="Invalid", index=False
    )

    flagged.drop(
        columns=["_reason_count"], errors="ignore"
    ).to_excel(
        writer, sheet_name="Anomalies", index=False
    )


# --------------------------------------------------
# 8. FORMAT EXCEL WORKBOOK
# --------------------------------------------------

wb = load_workbook(excel_path)

header_fill = PatternFill(
    "solid", fgColor="245C45"
)

high_fill = PatternFill(
    "solid", fgColor="FCE4D6"
)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if ws.max_row > 0 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions

    # Format headers
    for cell in ws[1]:
        cell.font = Font(
            bold=True, color="FFFFFF"
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[1].height = 30

    # Adjust column widths
    for column_cells in ws.columns:
        letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            (
                len(str(cell.value))
                if cell.value is not None
                else 0
            )
            for cell in column_cells[:500]
        )

        ws.column_dimensions[letter].width = min(
            max_length + 2, 35
        )

# Highlight High-severity rows in Summary
ws = wb["Summary"]

headers = {
    cell.value: cell.column
    for cell in ws[1]
}

if "severity" in headers and ws.max_row >= 2:
    severity_col = get_column_letter(
        headers["severity"]
    )

    ws.conditional_formatting.add(
        f"A2:{get_column_letter(ws.max_column)}{ws.max_row}",
        FormulaRule(
            formula=[
                f'${severity_col}2="High"'
            ],
            fill=high_fill,
        ),
    )

wb.save(excel_path)

print("Excel workbook created and formatted.")


# --------------------------------------------------
# 9. PREPARE REPORT TEXT
# --------------------------------------------------

date_min = df["transaction_date"].min()
date_max = df["transaction_date"].max()

date_min_text = (
    date_min.strftime("%Y-%m-%d")
    if pd.notna(date_min)
    else "Unavailable"
)

date_max_text = (
    date_max.strftime("%Y-%m-%d")
    if pd.notna(date_max)
    else "Unavailable"
)

run_time = (
    datetime.now() - start_time
).total_seconds()


# Format the dimension results as readable bullet points
dimension_lines = [
    f"- **{dimension}:** {count:,} failed check rows"
    for dimension, count in dimension_summary.items()
]

# Format top issues as readable bullet points
issue_lines = []

for _, row in top_issues.iterrows():
    issue_lines.append(
        f"- **{row['check_name']}** — "
        f"{row['severity']} severity; "
        f"{int(row['failed_rows']):,} failed rows "
        f"({row['failure_percent']:.2f}%)."
    )

# Format anomaly highlights without raw Markdown tables
anomaly_lines = []

for _, row in top_anomalies.iterrows():
    transaction_id = row.get(
        "transaction_id", "Unknown ID"
    )

    customer_id = row.get(
        "customer_id", "Unknown"
    )

    amount_value = row.get("amount", None)

    if pd.isna(amount_value):
        amount_text = "Missing"
    else:
        amount_text = f"{amount_value:,.2f}"

    currency = row.get("currency", "")

    if pd.isna(currency):
        currency = ""

    reason = row.get(
        "flag_reason", "Not specified"
    )

    anomaly_lines.extend([
        f"**Transaction: {transaction_id}**",
        f"- Customer: {customer_id}",
        f"- Amount: {currency} {amount_text}".strip(),
        f"- Detection reasons: {reason}",
        "",
    ])

if not anomaly_lines:
    anomaly_lines = [
        "No transactions were flagged by the anomaly detector."
    ]


# --------------------------------------------------
# 10. GENERATE MARKDOWN REPORT
# --------------------------------------------------

lines = [
    "# Transaction Data Quality Report",
    "",
    "## 1. Overview",
    "",
    f"- **Rows scanned:** {len(df):,}",
    f"- **Transaction date range:** {date_min_text} to {date_max_text}",
    f"- **Report generation time:** {run_time:.2f} seconds",
    "",
    "## 2. Overall Data Quality",
    "",
    f"- **Clean rows:** {clean_rows:,}",
    f"- **Total rows used for score:** {total_rows:,}",
    f"- **Data quality score:** {dq_score:.2f}%",
    f"- **Score calculation basis:** {score_basis}",
    "",
    "Note: the score reflects the stated calculation basis, "
    "not necessarily every quality check in the project.",
    "",
    "## 3. Results by Dimension",
    "",
    *dimension_lines,
    "",
    "The counts above are summed across checks. "
    "A single transaction may fail multiple checks.",
    "",
    "![Data quality failures by dimension](dq_dimension_chart.png)",
    "",
    "## 4. Top Data Quality Issues",
    "",
    *issue_lines,
    "",
    "## 5. Anomaly Highlights",
    "",
    *anomaly_lines,
    "",
    "The transactions above are flagged for review. "
    "An anomaly is not automatically a confirmed error or fraud.",
    "",
    "## 6. Recommendations",
    "",
    "- Require customer IDs when transactions are submitted.",
    "- Enforce transaction ID uniqueness where required by the business rules.",
    "- Validate email formats before saving records.",
    "- Validate transaction amounts and permitted currencies.",
    "- Investigate repeated high-value or unusual transactions.",
    "- Review flagged records before taking corrective action.",
    "",
    "## 7. Limitations",
    "",
    "- Quality rules and anomaly thresholds may require tuning.",
    "- Anomaly detection can produce false positives and false negatives.",
    "- Detector flags identify records for review, not confirmed fraud.",
    "- The overall score depends on the checks and labels used in its calculation.",
    "- The detail sheets use the specific rules implemented in this report script.",
]

markdown_path = REPORTS / "data_quality_report.md"

markdown_path.write_text(
    "\n".join(lines) + "\n",
    encoding="utf-8",
)


# --------------------------------------------------
# 11. FINAL OUTPUT
# --------------------------------------------------

print("\n" + "=" * 50)
print("PHASE 7 COMPLETED")
print("=" * 50)

print(f"Excel workbook : {excel_path}")
print(f"Markdown report: {markdown_path}")
print(f"Dimension chart: {chart_path}")

print(f"\nRows scanned: {len(df):,}")
print(f"Clean rows used in score: {clean_rows:,}")
print(f"Data quality score: {dq_score:.2f}%")
print(f"Anomaly rows flagged: {len(flagged):,}")
print(f"Execution time: {run_time:.2f} seconds")

